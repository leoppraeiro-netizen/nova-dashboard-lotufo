"""FRC gerada pela dashboard (2026-10-09): a FRC do consorcio do cadastro de exemplo (Obra Consorcio, 14a medicao
de execucao) reproduzida a partir das quatro notas, checklist que bloqueia, numeracao e a FRC assinada no envio.

As notas sao textos de fixture (uma da empresa propria pelo portal de Cuiaba, outra da consorciada Beta pelo
Centi) com numero, municipio, CNPJ e valores trocados pelos do espelho de exemplo (R$ 10.000.000,07 cheio:
Cuiaba 60% e Varzea Grande 40%, metade para cada empresa). INSS/ISS/IRRF seguem a regra do espelho (11% e 5%
sobre 40%, 1,2% sobre o total). O cadastro da FRC do consorcio de exemplo e' semeado pelo proprio teste: em
producao ele mora na migracao 032."""
import io
import json

import openpyxl
import pypdf
import pytest

from app import cadastro_aplicar, cadastro_arquivos, espelho, espelhos, frc, medicoes, notas, notas_envio, protocolo
from tests.conftest import CADASTRO_EXEMPLO, _pdf_com_texto
from tests.test_notas import CNPJ_BETA, CNPJ_PROPRIA, _fixture, _texto, conn, usuario  # noqa: F401  (fixtures)

LOTUFO, BETA = CNPJ_PROPRIA, CNPJ_BETA
CONSORCIO = "consorcio-exemplo"
ULTIMA = 40                      # a ultima FRC emitida antes do sistema: a primeira gerada sai 41
# (empresa, numero, municipio, bruto, inss, iss, irrf), em centavos. A ultima e' do leiaute de Cuiaba. O cheio
# (R$ 10.000.000,07) tem centavos: a metade de Varzea Grande cai em .5 e o ISS/INSS/IRRF truncam.
NOTAS_FRC = [
    ("beta", 20260000204, "Várzea Grande", 200000001, 8800000, 4000000, 2400000),
    ("beta", 20260000205, "Cuiabá", 300000002, 13200000, 6000000, 3600000),
    ("lotufo", 2284, "Cuiabá", 300000002, 13200000, 6000000, 3600000),
    ("lotufo", 2286, "Várzea Grande", 200000002, 8800000, 4000000, 2400000),
]
DADOS_FRC = {
    "logo": None,
    "prestador": {"razao_social": "CONSÓRCIO EXEMPLO", "endereco": "RUA DE EXEMPLO, 100, SALA 1, BAIRRO EXEMPLO,",
                  "municipio": "CUIABÁ - MT / BRASIL", "cep": "00000-000", "cnpj": "31.200.303/0001-81",
                  "inscricao_municipal": "123456", "telefone": "(00) 0000-0000", "email": "contato@exemplo.com.br"},
    "tomador": {"razao_social": "SECRETARIA DE EXEMPLO", "endereco": "AV. DE EXEMPLO", "municipio": "CUIABÁ - MT / BRASIL",
                "cep": "00000-001", "cnpj": "31.300.101/0001-01", "inscricao_municipal": "654321",
                "telefone": "(00) 0000-0001", "email": "contato@exemplo.gov.br"},
    "descricao": {
        "execucao": "{n}ª MEDIÇÃO DE CONTRATO - PERÍODO DE {inicio} A {fim} - OBRAS DE EXEMPLO, CONFORME IC 002/2025.",
        "reajuste": "REAJUSTE DA {n}ª MEDIÇÃO DE CONTRATO - PERÍODO DE {inicio} A {fim} - OBRAS DE EXEMPLO."},
    "rotulo_empresa": "RAZÃO SOCIAL",
    "consorciadas": [{"chave": "lotufo", "cnpj": "31.200.101/0001-30", "razao_social": "CONSTRUTORA EXEMPLO LTDA"},
                     {"chave": "beta", "cnpj": "31.200.202/0001-00", "razao_social": "ENGENHARIA BETA S.A."}],
}


def _semear_frc(conn, ultimo=ULTIMA, **dados):
    """O cadastro da FRC do Consorcio Exemplo (formulario A) gravado direto, sem logo, por cima do frc.yaml."""
    conn.execute("INSERT OR REPLACE INTO frc_consorcios (consorcio, formulario, ultimo_anterior, dados) "
                 "VALUES (?, 'A', ?, ?)", (CONSORCIO, ultimo, json.dumps({**DADOS_FRC, **dados}, ensure_ascii=False)))
    conn.commit()


def _cuiaba(cnpj, numero, municipio, bruto, inss, iss, irrf) -> str:
    f = espelho._fmt
    ret = inss + iss + irrf
    t = _fixture("3102-integra-vg.txt").replace(LOTUFO, cnpj)
    t = t.replace("Número da Nota Fiscal\n3102", f"Número da Nota Fiscal\n{numero}").replace("Várzea Grande", municipio)
    t = t.replace("12ª MEDIÇÃO", "14ª MEDIÇÃO").replace("PERÍODO DE 01/06/2026 A 30/06/2026", "PERÍODO DE 01/08/2026 A 31/08/2026")
    t = t.replace("R$846.788,31", "R$" + f(bruto)).replace("Vl. ISSQN: R$16.935,77", "Vl. ISSQN: R$" + f(iss))
    t = t.replace("Vl. IRRF: R$10.161,45", "Vl. IRRF: R$" + f(irrf)).replace("Vl. CP Retido: R$37.258,68", "Vl. CP Retido: R$" + f(inss))
    return t.replace("R$64.355,90 R$2.876,20 R$319,58 R$782.432,41 R$782.432,41",
                     f"R${f(ret)} R$2.876,20 R$319,58 R${f(bruto - ret)} R${f(bruto - ret)}")


def _centi(numero, municipio, bruto, inss, iss, irrf) -> str:
    f = espelho._fmt
    t = _fixture("2026812-integra-beta-vg.txt").replace("20260000812", str(numero))   # tambem na chave de acesso
    t = t.replace("VÁRZEA GRANDE - MT", municipio.upper() + " - MT").replace("11ª MEDIÇÃO", "14ª MEDIÇÃO")
    t = t.replace("PERÍODO DE 01/05/2026 A 31/05/2026", "PERÍODO DE 01/08/2026 A 31/08/2026")
    for de, para in (("914.477,97", bruto), ("40.237,02", inss), ("10.973,73", irrf), ("51.210,75", inss + irrf),
                     ("18.289,56", iss), ("844.977,66", bruto - inss - irrf - iss)):
        t = t.replace(de, f(para))
    return t


def _texto_da(nota) -> str:
    empresa, numero, municipio, *v = nota
    if empresa == "beta":
        return _centi(numero, municipio, *v)
    return _cuiaba(LOTUFO, numero, municipio, *v)


def _soma(i: int) -> int:
    return sum(n[i] for n in NOTAS_FRC)


@pytest.fixture
def consorcio14(conn, usuario):  # noqa: F811
    """Cadastro de exemplo, FRC do consorcio semeada, espelho da 14a com o cheio de R$ 10.000.000,07."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    _semear_frc(conn)
    return _medicao14(conn, usuario)


def _medicao14(conn, usuario):  # noqa: F811
    cid = conn.execute("SELECT id FROM contratos WHERE nome_curto='Obra Consórcio'").fetchone()[0]
    eid = espelhos.gerar(conn, usuario, cid, 14, "2026-08", {"medicao": _soma(3)})
    mid = conn.execute("SELECT m.id FROM espelho_medicoes em JOIN medicoes m ON m.id=em.medicao_id "
                       "WHERE em.espelho_id=? AND m.tipo='execucao'", (eid,)).fetchone()[0]
    medicoes.marcar_espelho_enviado(conn, 1, mid)
    return cid, mid


def _anexar(conn, usuario, mid, _texto, lista):  # noqa: F811
    out = []
    for i, nota in enumerate(lista):
        _texto(_texto_da(nota))
        out.append(notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto(f"nota {nota[1]}"), f"NFSE {i}.pdf"))
    return out


def test_frc_reproduzida_das_quatro_notas(conn, usuario, consorcio14, _texto):  # noqa: F811
    _, mid = consorcio14
    s = frc.situacao(conn, mid)
    assert s["disponivel"] and (s["ultima"], s["sugerido"]) == (ULTIMA, ULTIMA + 1) and s["frc"] is None
    assert [(i["empresa"], i["municipio"]) for i in s["itens"]] == [
        ("Exemplo", "Cuiabá"), ("Beta", "Cuiabá"), ("Exemplo", "Várzea Grande"), ("Beta", "Várzea Grande")]
    r = _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    assert [x["frc_gerada"] for x in r] == [None] * 3 + [ULTIMA + 1]      # a ultima nota gera sozinha, com ultima + 1
    s = frc.situacao(conn, mid)
    assert s["completa"] and s["bloqueio"] is None and s["frc"]["numero"] == ULTIMA + 1 and s["ultima"] == ULTIMA + 1
    assert all(i["situacao"] == "anexada" for i in s["itens"])
    assert {i["numero"] for i in s["itens"] if i["empresa"] == "Beta"} == {"2026204", "2026205"}
    f = conn.execute("SELECT * FROM frcs WHERE medicao_id=?", (mid,)).fetchone()
    bruto, inss, iss, irrf = _soma(3), _soma(4), _soma(5), _soma(6)
    liquido = bruto - inss - iss - irrf
    assert (f["bruto_centavos"], f["inss_centavos"], f["irrf_centavos"], f["iss_centavos"], f["liquido_centavos"]) == (
        1000000007, 44000000, 12000000, 20000000, 924000007) == (bruto, inss, irrf, iss, liquido)
    linhas = json.loads(f["notas"])
    assert [x["numero"] for x in linhas] == ["2284", "2286", "2026204", "2026205"]      # por consorciada, depois pelo numero
    assert [(x["bruto"], x["retencoes"], x["liquido"]) for x in linhas][-2:] == [(200000001, 15200000, 184800001),
                                                                                (300000002, 22800000, 277200002)]
    texto = frc.arquivo(conn, mid, "pdf").read_bytes()
    t = pypdf.PdfReader(io.BytesIO(texto)).pages[0].extract_text()
    for trecho in ("FATURA DE RECEBIMENTO DE CONSÓRCIO", f"{ULTIMA + 1:05d}", "GL.SGI.FRP.300", "CONSÓRCIO EXEMPLO",
                   "31.200.303/0001-81", "14ª MEDIÇÃO DE CONTRATO - PERÍODO DE 01/08/2026 A 31/08/2026",
                   "10.000.000,07", "440.000,00", "120.000,00", "200.000,00", "9.240.000,07", "2026204",
                   "ENGENHARIA BETA S.A.", "contato@exemplo.gov.br"):
        assert trecho in t, trecho
    wb = openpyxl.load_workbook(frc.arquivo(conn, mid, "xlsx"))
    valores = [c.value for linha in wb.active.iter_rows() for c in linha if c.value is not None]
    assert 9240000.07 in valores and f"{ULTIMA + 1:05d}" in valores and "2026204" in valores
    assert frc.nome_arquivo(conn, mid, "pdf") == f"FRC {ULTIMA + 1} - CONSÓRCIO EXEMPLO (14ª MEDIÇÃO).pdf"


def test_frc_do_cadastro_sai_com_o_logo_e_o_tomador_do_cadastro(conn, usuario, _texto):  # noqa: F811
    """Issue #110: sem semente da migracao, a FRC vem do frc.yaml (cadastro-exemplo). O logo, copiado para a pasta
    de dados pelo `cadastro aplicar`, vai no PDF e na planilha; o e-mail do tomador e' o do cadastro."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    _, mid = _medicao14(conn, usuario)
    assert _anexar(conn, usuario, mid, _texto, NOTAS_FRC)[-1]["frc_gerada"] == ULTIMA + 1
    pagina = pypdf.PdfReader(io.BytesIO(frc.arquivo(conn, mid, "pdf").read_bytes())).pages[0]
    assert len(pagina.images) == 1 and "fiscal@exemplo.gov.br" in pagina.extract_text()
    assert len(openpyxl.load_workbook(frc.arquivo(conn, mid, "xlsx")).active._images) == 1


def test_frc_com_logo_que_nao_esta_na_pasta_de_dados_e_recusada(conn, usuario, _texto):  # noqa: F811
    """Deploy sem `cadastro aplicar` (issue #110): a linha antiga aponta para o logo que saiu do app. A FRC nao sai
    sem o papel timbrado: a geracao automatica nao gera e a manual explica o que falta."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    dados = json.loads(conn.execute("SELECT dados FROM frc_consorcios").fetchone()[0])
    conn.execute("UPDATE frc_consorcios SET dados=?", (json.dumps({**dados, "logo": "frc-integra-brt.jpg"}),))
    conn.commit()
    _, mid = _medicao14(conn, usuario)
    assert _anexar(conn, usuario, mid, _texto, NOTAS_FRC)[-1]["frc_gerada"] is None
    with pytest.raises(ValueError, match="Logo da FRC não encontrado na pasta de dados: aplique o cadastro"):
        frc.gerar(conn, usuario, mid, ULTIMA + 1)
    assert conn.execute("SELECT COUNT(*) FROM frcs").fetchone()[0] == 0


def test_natureza_sem_descricao_no_cadastro_recusa_com_mensagem(conn, usuario, _texto):  # noqa: F811
    """O frc.yaml pode trazer so' as naturezas que o consorcio usa (Complexo Leblon nao tem projeto): medicao de
    natureza sem descricao nao gera, com a mensagem, sem estourar."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    _semear_frc(conn, descricao={"reajuste": "REAJUSTE DA {n}ª MEDIÇÃO - {inicio} A {fim}"})
    _, mid = _medicao14(conn, usuario)
    assert _anexar(conn, usuario, mid, _texto, NOTAS_FRC)[-1]["frc_gerada"] is None
    with pytest.raises(ValueError, match="Cadastro da FRC sem descrição para 14ª medição"):
        frc.gerar(conn, usuario, mid, ULTIMA + 1)


def test_checklist_com_nota_faltando_bloqueia(conn, usuario, consorcio14, _texto):  # noqa: F811
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC[:3])
    s = frc.situacao(conn, mid)
    assert not s["completa"] and s["frc"] is None
    assert s["bloqueio"] == "Falta 1 nota: Exemplo · Várzea Grande"
    with pytest.raises(ValueError, match="Exemplo"):
        frc.gerar(conn, usuario, mid, ULTIMA + 1)


def test_nota_sem_leitura_confiavel_bloqueia_e_rascunho_nao_conta(conn, usuario, consorcio14, _texto):  # noqa: F811
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC[:3])
    # rascunho da ultima nota: mesmo nome de arquivo, sem numero no conteudo -> nao atende a saida
    _texto(_texto_da(NOTAS_FRC[3]).replace("Número da Nota Fiscal\n2286", "Número da Nota Fiscal\n0"))
    notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("rascunho"), "NFSE 2286.pdf")
    s = frc.situacao(conn, mid)
    assert [x["motivo"] for x in s["nao_lidas"]] == ["sem-numero"] and not s["completa"]
    # a nota de verdade, com o liquido impresso que nao fecha com bruto - retencoes
    _texto(_texto_da(NOTAS_FRC[3]).replace("R$1.848.000,02 R$1.848.000,02", "R$1.848.000,03 R$1.848.000,03"))
    r = notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("2286 errada"), "NFSE 2286.pdf")
    assert r["frc_gerada"] is None
    s = frc.situacao(conn, mid)
    assert s["completa"] and s["frc"] is None
    assert s["bloqueio"].startswith("Nota sem leitura confiável — NF 2286 (Exemplo)")
    with pytest.raises(ValueError, match="sem leitura confiável"):
        frc.gerar(conn, usuario, mid, ULTIMA + 1)


def test_arrastar_de_novo_rele_a_nota_que_nao_tinha_sido_lida(conn, usuario, consorcio14, _texto):  # noqa: F811
    # 09/10/2026: as notas novas da consorciada chegaram antes do leitor conhecer o modelo; a pessoa arrasta o
    # MESMO PDF de novo depois da correcao e a nota passa a contar (e a FRC sai sozinha)
    _, mid = consorcio14
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC[:3])
    _texto("NOTA FISCAL DE SERVIÇOS ELETRÔNICA\nleiaute que o leitor ainda nao conhece")
    a = notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("nota 2286"), "NFSE 2286.pdf")
    assert not a["leitura"]["ok"] and a["frc_gerada"] is None and not frc.situacao(conn, mid)["completa"]
    _texto(_texto_da(NOTAS_FRC[3]))
    b = notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("nota 2286"), "NFSE 2286.pdf")
    assert b["repetido"] and b["documento_id"] == a["documento_id"]
    assert b["frc_gerada"] == ULTIMA + 1 and frc.situacao(conn, mid)["completa"]


def test_numero_sugerido_e_ultima_registrada(conn, usuario, consorcio14, _texto):  # noqa: F811
    cid, mid = consorcio14
    n = ULTIMA + 1
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    assert frc.ultima(conn, CONSORCIO) == n
    frc.gerar(conn, usuario, mid, n + 8)                  # numero digitado: gera de novo e a ultima acompanha
    s = frc.situacao(conn, mid)
    assert (s["frc"]["numero"], s["ultima"], s["sugerido"]) == (n + 8, n + 8, n + 8)
    assert conn.execute("SELECT COUNT(*) FROM frcs WHERE medicao_id=? AND substituida_em IS NOT NULL",
                        (mid,)).fetchone()[0] == 1
    frc.gerar(conn, usuario, mid, n)                      # volta ao n: a n + 8 foi substituida
    assert frc.ultima(conn, CONSORCIO) == n
    outra = medicoes.criar(conn, 1, {"contrato_id": cid, "numero": 15, "tipo": "execucao", "competencia": "2026-09"})
    assert frc.situacao(conn, outra)["sugerido"] == n + 1
    # o mesmo numero em outra medicao do consorcio e' recusado (aqui a outra nem tem as notas: o numero vem antes)
    conn.execute("INSERT INTO frcs (medicao_id, consorcio, numero, data_geracao, bruto_centavos, inss_centavos, "
                 "irrf_centavos, iss_centavos, liquido_centavos, notas, pdf, pdf_sha256, xlsx, xlsx_sha256) "
                 "VALUES (?, ?, ?, '2026-10-09', 0, 0, 0, 0, 0, '[]', 'x', 'x', 'x', 'x')", (outra, CONSORCIO, n + 1))
    conn.commit()
    with pytest.raises(ValueError, match=f"FRC nº {n + 1} já é a da 15ª Medição"):
        frc.gerar(conn, usuario, mid, n + 1)


def test_frc_gerada_nao_satisfaz_o_envio_e_a_assinada_sim(conn, usuario, consorcio14, _texto):  # noqa: F811
    cid, mid = consorcio14
    n = ULTIMA + 1
    _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    c = protocolo._contrato(conn, cid)
    m = dict(conn.execute("SELECT * FROM medicoes WHERE id=?", (mid,)).fetchone())
    assert frc.situacao(conn, mid)["frc"] is not None
    assert protocolo.faltas_consorcio(conn, c, [m]) == ["14ª Medição: FRC assinada não anexada"]
    itens, _, _ = notas_envio._itens(conn, c, "sinfra", [m], __import__("datetime").date(2026, 10, 9))
    item = next(i for i in itens if i["chave"] == f"frc:{mid}")
    assert (item["estado"], item["acao"]) == ("falta", "anexar_frc")
    # issue #148: a etapa 4 diz que falta a assinada e oferece baixar a gerada
    assert item["frc_gerada"] == n
    assert item["detalhe"] == "Falta a assinada: baixe a FRC gerada, assine no DocuSign e anexe aqui"
    assert next(i for i in itens if i["chave"] == "consorcio")["estado"] == "ok"
    _texto("qualquer coisa")
    notas.anexar_pdf(conn, usuario, mid, _pdf_com_texto("assinada"), f"FRC {n} assinada.pdf", tipo_doc="fatura")
    assert protocolo.faltas_consorcio(conn, c, [m]) == []
    itens, _, _ = notas_envio._itens(conn, c, "sinfra", [m], __import__("datetime").date(2026, 10, 9))
    item = next(i for i in itens if i["chave"] == f"frc:{mid}")
    assert (item["estado"], item["detalhe"]) == ("ok", f"Assinada anexada · FRC {n} assinada.pdf")
    assert frc.situacao(conn, mid)["assinada"] is True


def test_consorcio_sem_formulario_e_obra_propria(conn, usuario):  # noqa: F811
    cad = cadastro_arquivos.preparar(CADASTRO_EXEMPLO)
    cadastro_aplicar.aplicar(conn, 1, {**cad, "frc": []}, ensaio=False)          # consorcio sem FRC no cadastro
    consorcio = conn.execute("SELECT id FROM contratos WHERE nome_curto='Obra Consórcio'").fetchone()[0]
    mid = medicoes.criar(conn, 1, {"contrato_id": consorcio, "numero": 50, "tipo": "execucao", "competencia": "2026-09"})
    assert frc.situacao(conn, mid) == {"disponivel": False, "consorcio": "Consórcio Exemplo", "motivo": frc.INDISPONIVEL}
    with pytest.raises(ValueError, match="ainda não disponível"):
        frc.gerar(conn, usuario, mid, 1)
    propria = conn.execute("SELECT id FROM contratos WHERE nome_curto='Obra Própria'").fetchone()[0]
    mid = medicoes.criar(conn, 1, {"contrato_id": propria, "numero": 50, "tipo": "execucao", "competencia": "2026-09"})
    assert frc.situacao(conn, mid)["disponivel"] is False and frc.situacao(conn, mid)["motivo"] is None


def test_rotas_da_frc(client_admin):
    """Obra propria: a FRC nao se aplica, baixar e' 404 e gerar e' 422 com o motivo."""
    from tests.conftest import novo_contrato
    cid = novo_contrato({"orgao": "SINFRA", "numero_instrumento": "IC 1/2026", "objeto": "Obra",
                         "valor_original": "100.000,00"})
    mid = client_admin.post(f"/api/contratos/{cid}/medicoes",
                            json={"numero": 1, "tipo": "execucao", "competencia": "2026-08"}).json()["id"]
    r = client_admin.get(f"/api/medicoes/{mid}/frc")
    assert r.status_code == 200 and r.json()["disponivel"] is False
    assert client_admin.get(f"/api/medicoes/{mid}/frc/pdf").status_code == 404
    assert client_admin.get(f"/api/medicoes/{mid}/frc/doc").status_code == 404
    r = client_admin.post(f"/api/medicoes/{mid}/frc", json={"numero": 22})
    assert r.status_code == 422 and "consórcio" in r.json()["detail"]
    assert client_admin.post(f"/api/medicoes/{mid}/frc", json={"numero": 0}).status_code == 422
    assert client_admin.get("/api/medicoes/99999/frc").status_code == 404


def test_numero_digitado_abaixo_da_semente_reproduz_a_frc_ja_emitida(conn, usuario, _texto):  # noqa: F811
    """Consorcio cuja ultima FRC antes do sistema e' alta (51) e cujo formulario tem outro rotulo e rodape: a
    gerada sozinha sai 52, e o numero digitado 50 (uma FRC emitida fora de ordem) a regera com o texto certo."""
    cadastro_aplicar.aplicar(conn, 1, cadastro_arquivos.preparar(CADASTRO_EXEMPLO), ensaio=False)
    _semear_frc(conn, ultimo=51, rotulo_empresa="EMPRESA CONSORCIADA",
                descricao={"execucao": "{n}° MEDIÇÃO DA OBRA DE EXEMPLO, PERÍODO {inicio} A {fim}. DADOS BANCÁRIOS: "
                                       "BANCO EXEMPLO AGENCIA: 0001-0 CONTA CORRENTE: 00000-0"})
    cid = conn.execute("SELECT id FROM contratos WHERE nome_curto='Obra Consórcio'").fetchone()[0]
    eid = espelhos.gerar(conn, usuario, cid, 28, "2026-08", {"medicao": _soma(3)})
    mid = conn.execute("SELECT m.id FROM espelho_medicoes em JOIN medicoes m ON m.id=em.medicao_id "
                       "WHERE em.espelho_id=? AND m.tipo='execucao'", (eid,)).fetchone()[0]
    medicoes.marcar_espelho_enviado(conn, 1, mid)
    r = _anexar(conn, usuario, mid, _texto, NOTAS_FRC)
    assert r[-1]["frc_gerada"] == 52
    frc.gerar(conn, usuario, mid, 50)
    f = conn.execute("SELECT * FROM frcs WHERE medicao_id=? AND substituida_em IS NULL", (mid,)).fetchone()
    assert (f["numero"], f["bruto_centavos"], f["liquido_centavos"]) == (50, 1000000007, 924000007)
    t = " ".join(pypdf.PdfReader(io.BytesIO(frc.arquivo(conn, mid, "pdf").read_bytes())).pages[0].extract_text().split())
    for trecho in ("00050", "CONSÓRCIO EXEMPLO", "EMPRESA CONSORCIADA", "ENGENHARIA BETA S.A.",
                   "28° MEDIÇÃO DA OBRA DE EXEMPLO, PERÍODO 01/08/2026 A 31/08/2026", "AGENCIA: 0001-0", "9.240.000,07"):
        assert trecho in t, trecho
