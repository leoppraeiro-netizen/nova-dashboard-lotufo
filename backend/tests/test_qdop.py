"""QDOP (GL.SGI.FRP.DE.250) exportado no leiaute oficial (bloco 8)."""
import datetime
import io

import openpyxl

from app import qdop
from tests.test_diretoria import cenario  # noqa: F401  (fixture)

D = datetime.date
HOJE = D(2026, 9, 24)


def test_linha_na_cota_com_faturados_recebidos(cenario):  # noqa: F811
    c, ids = cenario          # consorcio 50%: original 10.000.000,00; aditivo 2.000.000,00; recebido exec 4.000.000,00; reaj 100.000,00
    linha = next(x for x in qdop.linhas(c, HOJE) if x["contrato_id"] == ids["consorcio"])
    assert (linha["empreita"], linha["direito"]) == ("CONSÓRCIO", 0.5)
    assert (linha["pi"], linha["va"], linha["pi_va"]) == (500_000_000, 100_000_000, 600_000_000)
    assert (linha["faturados"], linha["reajuste_recebido"], linha["saldo"]) == (200_000_000, 5_000_000, 400_000_000)


def test_status_paralisada_e_concluida(cenario):  # noqa: F811
    c, ids = cenario
    status = lambda: {x["contrato_id"]: x["status"] for x in qdop.linhas(c, HOJE)}  # noqa: E731
    assert status()[ids["propria"]] == "EM ANDAMENTO"
    c.execute("INSERT INTO instrumentos (contrato_id, chave, tipo, data, dados) VALUES (?, 'par-1', 'paralisacao', "
              "'2026-08-01', '{}')", (ids["propria"],))
    c.execute("UPDATE contratos SET encerramento_iniciado_em='2026-09-01' WHERE id=?", (ids["consorcio"],)); c.commit()
    assert (status()[ids["propria"]], status()[ids["consorcio"]]) == ("PARALISADA", "CONCLUÍDA")


def test_xlsx_no_leiaute_oficial(cenario):  # noqa: F811
    c, ids = cenario
    ws = openpyxl.load_workbook(io.BytesIO(qdop.gerar_xlsx(c, HOJE))).active
    linhas = qdop.linhas(c, HOJE)
    assert ws["S1"].value == "FRM.EXEMPLO.001"   # codigo do formulario no modelo de cadastro-exemplo/
    assert ws["A13"].value == "SEQ." and ws["S13"].value == "DESCRIÇÃO DA OBRA"
    assert ws["A14"].value == 1 and ws["I14"].value == 5_000_000.00        # PI do consorcio (inicio 2024) em reais
    ultima = 13 + len(linhas)
    assert ws.cell(ultima + 2, 2).value == "DEPARTAMENTO DE ENGENHARIA"
    assert ws["K11"].value == sum(x["pi_va"] for x in linhas) / 100
    assert ws.print_area.endswith(f"$S${ultima + 3}")


def test_nota_emitida_conta_como_faturado(cenario):  # noqa: F811
    c, ids = cenario
    c.execute("INSERT INTO medicoes (contrato_id, numero, tipo, competencia, estado, valor_protocolado_centavos) "
              "VALUES (?, 3, 'execucao', '2026-09', 'nota_emitida', 1000000)", (ids["consorcio"],)); c.commit()
    linha = next(x for x in qdop.linhas(c, HOJE) if x["contrato_id"] == ids["consorcio"])
    assert linha["faturados"] == 200_000_000 + 500_000
